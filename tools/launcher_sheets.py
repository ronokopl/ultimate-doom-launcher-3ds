"""One-shot: turn the sheets copied from Brutal Doom 3DS into the Doom Launcher 3DS source of truth.

Edits build.xlsx / platform.xlsx in place (identity, SD paths, generic runtime features) and adds the
Launcher (UI struts) and ModCache (on-device conversion rules) sheets. Re-running is harmless.
Status cells are left empty: tools/verify_build.py sets them from objective checks.
"""
import os
from openpyxl import load_workbook
from pk3 import SHEETS

BASE = 'sdmc:/3ds/DoomLauncher/'


def table(ws):
    hdr = [c.value for c in ws[1]]
    return hdr, {str(r[0].value): r for r in ws.iter_rows(min_row=2) if r[0].value is not None}


def set_row(ws, key, **vals):
    hdr, idx = table(ws)
    row = idx.get(key)
    if row is None:
        ws.append([key] + [None] * (len(hdr) - 1))
        row = list(ws.iter_rows(min_row=ws.max_row))[0]
    for k, v in vals.items():
        row[hdr.index(k)].value = v
    if 'status' in hdr:
        row[hdr.index('status')].value = None


def delete_rows(ws, keys):
    for r in reversed(range(2, ws.max_row + 1)):
        if str(ws.cell(r, 1).value) in keys:
            ws.delete_rows(r)


def new_sheet(wb, name, header, rows):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(header)
    for r in rows:
        ws.append([r.get(h) for h in header])


def build():
    p = os.path.join(SHEETS, 'build.xlsx')
    wb = load_workbook(p)
    ws = wb['Build']
    set_row(ws, 'title', value='Ultimate Doom Launcher', notes='SPEC: mod launcher CIA (name from the banner)')
    set_row(ws, 'short_title', value='Ultimate Doom Launcher', notes='SMDH')
    set_row(ws, 'publisher', value='ZDoom 2.8.1 / 3DS port', notes='SMDH')
    set_row(ws, 'unique_id', value='0xFDA3D', notes='homebrew range 0xF0000-0xFFFFF (auto; differs from Brutal Doom 3DS 0xFBD3D)')
    set_row(ws, 'product_code', value='CTR-H-DL3D', notes='makerom product code')
    set_row(ws, 'icon_source', value='doom2:STFST01', notes='SMDH icon: IWAD status-bar face (tools/gen_icon.py)')
    set_row(ws, 'max_debris', value='unlimited', notes='engine default: mods decide their own debris')
    # HOME Menu banner sound: the format of a banner known to play on hardware (Legend of Doom 3DS: PCM16 stereo
    # 32000 Hz, 2.95 s); the mono 44.1 kHz / 22.05 kHz banners stayed silent
    set_row(ws, 'banner_audio_rate', value='32000', notes='banner sound sample rate (Hz), as the hardware-tested banner')
    set_row(ws, 'banner_audio_s', value='2.9', notes='banner sound length (s); the HOME Menu limit is 3 s')
    set_row(ws, 'banner_audio_channels', value='2', notes='banner sound channels: stereo like the hardware-tested banner (mono did not play)')
    set_row(ws, 'gz_unique_id', value='0xFDA3E', notes='experimental GZDoom 4.7.1 engine CIA (../gzengine, udl/package.sh) for ZScript mods; Launcher btn_engine chainloads 00040000<id>00')
    set_row(ws, 'uz_unique_id', value='0xFDA3F', notes='experimental UZDoom 5.0.3 engine CIA (../uzengine, platform/3ds/udl/package.sh) for newer ZScript mods; Launcher btn_engine chainloads 00040000<id>00')
    set_row(ws, 'cache_deflate', value='1', notes='mod cache entries: 1 = deflated when >= 10% smaller (on the writer core, hidden behind image conversion; stored entries made caches ~7x bigger), 0 = stored')
    set_row(ws, 'cache_version', value='3', notes='ModCache format version; bump to invalidate every cache pk3 on the SD card')
    set_row(ws, 'ld_extra', value='-Wl,--no-gc-sections -Wl,--wrap=inflate -Wl,--wrap=ov_read -Wl,--wrap=fread -Wl,--wrap=fopen -Wl,--wrap=fseek -Wl,--wrap=svcSendSyncRequest',
            notes='keep ZDoom autoseg sections; --wrap hooks for the hitch profiler (profile_3ds.cpp), the stdio buffer (io_3ds.cpp) and the HOME handover IPC trace (i_input_3ds.cpp)')
    delete_rows(ws, {'sprite_png'})
    for unit, sub in (('platform/3ds/launcher_3ds.cpp', 'launcher'), ('platform/3ds/modcache_3ds.cpp', 'modcache')):
        set_row(wb['Sources'], unit, module='platform', compile='yes', reason='3DS platform layer (Subsystems sheet: %s)' % sub, flags='~')
    hdr, mods = table(wb['Modules'])
    for key in ('engine', 'platform'):	# lemon/re2c/gitinfo outputs are committed in engine/gensrc (were build/gensrc)
        cell = mods[key][hdr.index('includes')]
        cell.value = cell.value.replace('build/gensrc', 'engine/gensrc')
        mods[key][hdr.index('status')].value = None
    delete_rows(wb['Waivers'], {r[0].value for r in wb['Waivers'].iter_rows(min_row=2)})
    ws = wb['Patches']
    set_row(ws, 'engine/zdoom/src/v_video.cpp', change='__3DS__ GetHitlist calls I3DS_PrecacheSprites(spritelist)',
            reason='mods load death/gib frames on first use (in-game freezes): precache_3ds.cpp marks every state sprite of each monster class present')
    set_row(ws, 'engine/zdoom/src/s_sound.cpp', change='__3DS__ S_PrecacheLevel calls I3DS_PrecacheSounds()',
            reason='hook kept for the generic precache (sounds of monster states are left to the engine)')
    delete_rows(ws, {'engine/zdoom/src/thingdef/thingdef_function.cpp + ../namedef.h',
                     'engine/zdoom/src/thingdef/thingdef_properties.cpp + ../actor.h + ../p_enemy.cpp'})
    for f, change in COMPAT_PATCHES:
        set_row(ws, f, change=change, reason='Compat sheet: GZDoom 2.x DECORATE used by mods (first seen in m12-hdoom_techdemo11.pk3)')
    set_row(ws, 'engine/zdoom/src/textures/pngtexture.cpp', change='__3DS__ FPNGTexture::MakeTexture reads I3DS_PngMemory (thread-local lump bytes) when set',
            reason='ModCache: the writer core converts PNG lumps from memory, no WAD access off the main thread (modcache_3ds.cpp)')
    set_row(ws, 'engine/zdoom/src/textures/texturemanager.cpp', change='__3DS__ PrecacheLevel texture loop calls I3DS_PumpApt()',
            reason='long level loads (precaching every monster frame) still answer HOME / sleep / close (i_input_3ds.cpp)')
    set_row(ws, 'engine/zdoom/src/d_main.cpp', change='__3DS__ D_DoomMain calls I3DS_ModCacheBuild() after R_Init',
            reason='ModCache sheet: on first launch of a mod, convert its sprites/sounds with the engine\'s own decoders into a cache pk3')
    wb.save(p)


LAUNCHER = [
    # element, kind, screen, x, y, w, h, label, action, key, notes
    dict(element='title', kind='text', screen='top', x=0, y=0, w=400, h=24, label='Ultimate Doom Launcher', action='~', key='~', notes='title bar'),
    dict(element='credit', kind='text', screen='top', x=250, y=5, w=144, h=16, label='made by ronokopl', action='~', key='~', notes='author credit, right side of the title bar'),
    dict(element='summary', kind='text', screen='top', x=0, y=26, w=400, h=190, label='IWAD + load order + cache status', action='~', key='~', notes='what Launch will start'),
    dict(element='hints', kind='text', screen='top', x=0, y=218, w=400, h=22, label='button hints', action='~', key='~', notes='context help'),
    dict(element='tab_iwad', kind='tab', screen='bottom', x=0, y=0, w=79, h=26, label='IWAD', action='tab', key='L', notes='L/R cycle tabs'),
    dict(element='tab_mods', kind='tab', screen='bottom', x=80, y=0, w=79, h=26, label='Mods', action='tab', key='R', notes='L/R cycle tabs'),
    dict(element='tab_profiles', kind='tab', screen='bottom', x=160, y=0, w=79, h=26, label='Profiles', action='tab', key='~', notes='touch, or L/R'),
    dict(element='tab_options', kind='tab', screen='bottom', x=240, y=0, w=80, h=26, label='Options', action='tab', key='~',
         notes='touch, or L/R: global settings (settings.ini): mod caching on/off (off = no conversion, -nomodcache, caches not loaded), delete cached files'),
    dict(element='list', kind='list', screen='bottom', x=0, y=28, w=320, h=176, label='files / profiles', action='select', key='A', notes='8 rows of 22 px; D-pad moves, A selects/toggles, touch drag scrolls'),
    dict(element='btn_up', kind='button', screen='bottom', x=0, y=206, w=52, h=34, label='Up', action='move_up', key='X', notes='Mods tab: earlier in load order'),
    dict(element='btn_down', kind='button', screen='bottom', x=54, y=206, w=52, h=34, label='Down', action='move_down', key='Y', notes='Mods tab: later in load order'),
    dict(element='btn_save', kind='button', screen='bottom', x=108, y=206, w=52, h=34, label='Save', action='save_profile', key='SELECT', notes='name via software keyboard'),
    dict(element='btn_delete', kind='button', screen='bottom', x=162, y=206, w=40, h=34, label='Del', action='delete_profile', key='B', notes='Profiles tab only'),
    dict(element='btn_engine', kind='button', screen='bottom', x=204, y=206, w=54, h=34, label='Auto|ZDoom|GZDoom|UZDoom', action='engine', key='ZR',
         notes='cycles the engine, saved per profile (engine=): Auto = experimental GZDoom when a ticked mod is ZScript-only, else ZDoom 2.8.1; GZDoom / UZDoom = the separate UDL engine CIAs (gz_unique_id / uz_unique_id), started by chainload with gz_launch / uz_launch'),
    dict(element='btn_launch', kind='button', screen='bottom', x=260, y=206, w=60, h=34, label='Launch', action='launch', key='START', notes='saves the selection as the last profile; with ZDoom warns once when a ticked mod is ZScript-only (I3DS_ModScriptKind), START again launches anyway'),
]

COMPAT_PATCHES = [
    ('engine/zdoom/src/thingdef/thingdef_states.cpp', '__3DS__ a non-string state argument parses as an expression -> FxRuntimeStateIndex'),
    ('engine/zdoom/src/thingdef/thingdef_exp.h + thingdef_expression.cpp', '__3DS__ FxRuntimeStateIndex: jump offset computed at runtime'),
    ('engine/zdoom/src/thingdef/thingdef_function.cpp + engine/zdoom/src/namedef.h', '__3DS__ GetCVar("name") expression function'),
    ('engine/zdoom/src/menu/menudef.cpp', '__3DS__ MENUDEF AddOptionMenu, ScaleSlider, Slider gray-check argument'),
    ('engine/zdoom/src/thingdef/thingdef_properties.cpp + engine/zdoom/src/actor.h + engine/zdoom/src/p_enemy.cpp', '__3DS__ MissileChanceMult property (class meta, used by SuggestMissileAttack)'),
]

COMPAT = [
    # feature, kind, resolution, files, tested_with, notes
    dict(feature='computed state jump', kind='state argument', resolution='implemented',
         files='thingdef_states.cpp thingdef_exp.h thingdef_expression.cpp', tested_with='m12-hdoom_techdemo11.pk3',
         notes='A_JumpIf(x, CallACS("...")): value = offset from the calling state, 0 = no jump'),
    dict(feature='GetCVar', kind='expression function', resolution='implemented', files='thingdef_function.cpp namedef.h',
         tested_with='m12-hdoom_techdemo11.pk3', notes='float value of an engine or CVARINFO cvar; 0 when unknown'),
    dict(feature='MissileChanceMult', kind='actor property', resolution='implemented',
         files='thingdef_properties.cpp actor.h p_enemy.cpp', tested_with='m12-hdoom_techdemo11.pk3',
         notes='class meta value (no AActor field / savegame change), multiplies the missile-chance distance'),
    dict(feature='AddOptionMenu', kind='MENUDEF keyword', resolution='implemented', files='menu/menudef.cpp',
         tested_with='m12-hdoom_techdemo11.pk3', notes='appends items to an existing option menu; unknown menus are parsed and dropped'),
    dict(feature='Slider gray-check', kind='MENUDEF argument', resolution='implemented', files='menu/menudef.cpp',
         tested_with='m12-hdoom_techdemo11.pk3', notes='7th Slider argument parsed; the slider is never grayed out'),
    dict(feature='ScaleSlider', kind='MENUDEF item', resolution='implemented', files='menu/menudef.cpp',
         tested_with='m12-hdoom_techdemo11.pk3', notes='integer slider; the text shown for 0 / -1 is dropped'),
    dict(feature='CallACS', kind='expression function', resolution='already in 2.8.1', files='~',
         tested_with='m12-hdoom_techdemo11.pk3', notes='alias of ACS_NamedExecuteWithResult'),
    dict(feature='CLIENTSIDEONLY', kind='actor flag', resolution='already in 2.8.1', files='~',
         tested_with='m12-hdoom_techdemo11.pk3', notes='dummy flag (netcode only)'),
    dict(feature='WEAPON.CHEATNOTWEAPON', kind='weapon flag', resolution='already in 2.8.1', files='~',
         tested_with='m12-hdoom_techdemo11.pk3', notes='WIF_CHEATNOTWEAPON'),
]

MODCACHE = [
    # rule, source, detect, convert_to, cache_entry, notes
    dict(rule='image_png', source='lumps in the sprites, textures, patches and graphics namespaces of a selected mod', detect=r'\x89PNG',
         convert_to='doom patch (engine FTexture pixels + spans)', cache_entry='<namespace folder>/<LUMP>.lmp',
         notes='same lump name and namespace, loaded right after the mod so it overrides only that mod copy; flats stay PNG'),
    dict(rule='image_jpeg', source='lumps in the sprites, textures, patches and graphics namespaces of a selected mod', detect=r'\xff\xd8\xff',
         convert_to='doom patch (engine FTexture pixels + spans)', cache_entry='<namespace folder>/<LUMP>.lmp',
         notes='JPEG images decode even slower than PNG'),
    dict(rule='sfx_ogg', source='S_sfx entries whose lump is in a selected mod', detect='OggS', convert_to='wav pcm16 mono (rate halved when >= 32 kHz)',
         cache_entry='sounds/<LUMP>.wav', notes='same format ndspsound.cpp makes at runtime; music lumps are not sfx and stay OGG'),
    dict(rule='cache_key', source='mod file', detect='size + mtime + IWAD PLAYPAL crc + cache_version', convert_to='cache file name',
         cache_entry='cache/<mod name>-<key>.pk3', notes='a changed mod, IWAD palette or cache format rebuilds the cache; entries stored or deflated per Build cache_deflate; older keys of the same mod are deleted'),
    dict(rule='cache_skip', source='mod file', detect='nothing to convert', convert_to='marker',
         cache_entry='cache/<mod name>-<key>.none', notes='mods without PNG/JPEG images or OGG sfx are not rescanned'),
]


def platform():
    p = os.path.join(SHEETS, 'platform.xlsx')
    wb = load_workbook(p)
    ws = wb['Filesystem']
    delete_rows(ws, {'iwad_doom2', 'iwad_doom', 'mod_pk3', 'iwad_picker'})
    set_row(ws, 'config', path=BASE + 'zdoom.ini', source='runtime')
    set_row(ws, 'saves', path=BASE + 'save/', source='runtime')
    set_row(ws, 'args_file', path=BASE + 'args.txt', source='runtime (optional extra command-line words, e.g. +logfile)')
    set_row(ws, 'iwad_dir', path=BASE + 'iwads/', source='user: copy IWADs here (doom.wad, doom2.wad, plutonia.wad, tnt.wad, heretic.wad, hexen.wad, freedoom*.wad ...)')
    set_row(ws, 'mod_dir', path=BASE + 'mods/', source='user: copy mods here (.pk3 .wad .zip .pk7 .deh .bex)')
    set_row(ws, 'cache_dir', path=BASE + 'cache/', source='runtime (ModCache sheet)')
    set_row(ws, 'profiles', path=BASE + 'profiles.ini', source='runtime (launcher profiles; [last] = previous launch)')
    set_row(ws, 'app_dir', path=BASE, source='runtime (diagnostics: ndsp.txt, hitches.txt, fb_*.raw dumps)')
    set_row(ws, 'settings', path=BASE + 'settings.ini', source='runtime (launcher Options tab: global settings, e.g. modcache=1)')
    set_row(ws, 'gz_launch', path=BASE + 'gzdoom/launch.txt', source='runtime (launcher -> GZDoom engine command line, one word per line; the engine keeps gzdoom.ini and saves in gzdoom/)')
    set_row(ws, 'uz_launch', path=BASE + 'uzdoom/launch.txt', source='runtime (launcher -> UZDoom engine command line, one word per line; the engine keeps uzdoom.ini and saves in uzdoom/)')
    set_row(ws, 'autolaunch', path=BASE + 'autolaunch.txt', source='user, optional: a profile name (or last) to start without the launcher screens')
    ws = wb['Subsystems']
    delete_rows(ws, {'gore_precache'})
    set_row(ws, 'launcher', replaces='~', impl='platform/3ds/launcher_3ds.cpp', api='citro2d + swkbd',
            notes='Launcher sheet: pick IWAD, mods, load order, profiles; builds the engine command line')
    set_row(ws, 'modcache', replaces='~', impl='platform/3ds/modcache_3ds.cpp', api='engine',
            notes='ModCache sheet: first launch converts a mod\'s PNG/JPEG sprites and OGG sfx into cache/<mod>-<key>.pk3')
    set_row(ws, 'precache', replaces='~', impl='platform/3ds/precache_3ds.cpp', api='engine',
            notes='level start: every state sprite of each monster class present (cvar cl_3dsprecache)')
    set_row(ws, 'main', replaces='src/posix/sdl/i_main.cpp', impl='platform/3ds/i_main_3ds.cpp', api='libctru',
            notes='romfsInit, osSetSpeedupEnable, launcher, argv from the selection; quitting relaunches the launcher (CIA)')
    ws = wb['Controls']	# generic defaults (Brutal Doom aliases like kickem/grenadetoss do not exist in other mods)
    for button, command, notes in (('L', '+speed', 'run while held'), ('B', '+crouch', 'crouch while held'),
                                   ('SELECT', 'togglemap', 'full-screen automap on the top screen'),
                                   ('DUP', 'invuse', 'use inventory item (Heretic/Hexen/Strife, mods)'),
                                   ('DDOWN', 'centerview', 'level the view'),
                                   ('DLEFT', 'invprev', 'previous inventory item'), ('DRIGHT', 'invnext', 'next inventory item')):
        set_row(ws, button, command=command, notes=notes)
    ws = wb['Menu']
    set_row(ws, 'Precache monster frames', menu='3DS Options', cvar='cl_3dsprecache', values='Off|On',
            default='On', widget='option')
    ws = wb['Audio']
    set_row(ws, 'sfx_ogg', source_format='ogg', target='pcm16 mono (rate halved when >= 32 kHz)',
            where='runtime cache (modcache_3ds.cpp, first launch of a mod)', notes='ModCache sfx_ogg: decoded once on the 3DS, then read as WAV')
    set_row(ws, 'music_ogg', source_format='ogg', target='ogg stream', where='runtime stream', notes='music stays OGG (streamed)')
    new_sheet(wb, 'Launcher', ['element', 'kind', 'screen', 'x', 'y', 'w', 'h', 'label', 'action', 'key', 'notes', 'status'], LAUNCHER)
    new_sheet(wb, 'Compat', ['feature', 'kind', 'resolution', 'files', 'tested_with', 'notes', 'status'], COMPAT)
    new_sheet(wb, 'ModCache', ['rule', 'source', 'detect', 'convert_to', 'cache_entry', 'notes', 'status'], MODCACHE)
    wb.save(p)


if __name__ == '__main__':
    build()
    platform()
    print('sheets updated')
