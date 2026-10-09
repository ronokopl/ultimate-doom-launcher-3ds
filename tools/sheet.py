"""xlsx <-> rows helpers. Cell convention: '~' = intentionally none/inherit (checked); empty = unchecked."""
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.cell import WriteOnlyCell

HDR_FONT, HDR_FILL = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='404040')

def cell(v):
    if v is None or v == '':
        return '~'
    if isinstance(v, (list, tuple)):
        return ' '.join(map(str, v)) if v else '~'
    if isinstance(v, bool):
        return 'yes' if v else 'no'
    return v

def write(path, sheets):
    """sheets: {name: (columns, rows[dict])}. Every cell gets a value ('~' if absent)."""
    wb = Workbook(write_only=True)
    for name, (cols, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.freeze_panes = 'B2'
        hdr = []
        for c in cols:
            h = WriteOnlyCell(ws, value=c); h.font, h.fill = HDR_FONT, HDR_FILL; hdr.append(h)
        ws.append(hdr)
        for r in rows:
            ws.append([cell(r.get(c)) for c in cols])
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)

def read(path):
    """{sheet: [ {col: value-or-None} ]}; None marks an unchecked (empty) cell."""
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        cols = [c for c in next(it, ()) if c is not None]
        out[ws.title] = [dict(zip(cols, (None if v in (None, '') else v for v in row)))
                         for row in it if any(v not in (None, '') for v in row)]
    return out

def columns(rows, first):
    """first + every other key seen, in first-seen order."""
    seen = list(first)
    for r in rows:
        for k in r:
            if k not in seen:
                seen.append(k)
    return seen
