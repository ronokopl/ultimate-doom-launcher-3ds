"""After a successful build: set each build/platform row's `status` from an objective check (never by hand).

done  = the row's strut exists in the generated output AND the code consuming it compiled.
todo  = check failed (preflight then reports the row as unfinished).
"""
import os, re, struct, subprocess, zipfile
from openpyxl import load_workbook
from export import rows
from pk3 import ROOT, SHEETS

HEADER = os.path.join(ROOT, 'gen', 'c', 'bd3ds_sheets.h')
MAKEFILE = os.path.join(ROOT, 'build', 'Makefile')
ELF = os.path.join(ROOT, 'build', 'doomlauncher3ds.elf')
CIA = os.path.join(ROOT, 'build', 'doomlauncher3ds.cia')
CIA_KEYS = ('cia_version', 'system_mode', 'banner_image', 'banner_audio', 'smdh_flags', 'smdh_regions',
            'makerom_exe', 'bannertool_exe')
ROMFS = os.path.join(ROOT, 'build', 'romfs')
SIZE = r'C:\devkitPro\devkitARM\bin\arm-none-eabi-size.exe'

def obj_of(unit):
    return os.path.join(ROOT, 'build', 'obj', os.path.splitext(unit)[0] + '.o')

def built(unit):
    src, obj = os.path.join(ROOT, unit), obj_of(unit)
    return os.path.exists(obj) and os.path.exists(src) and os.path.getmtime(obj) >= os.path.getmtime(src)

def main():
    header = open(HEADER).read()
    makefile = open(MAKEFILE).read()
    sources = rows('build', 'Sources')
    compiled = {s['unit'] for s in sources if s.get('compile') == 'yes' and built(s['unit'])}
    platform_src = ''.join(open(os.path.join(ROOT, u), encoding='utf-8', errors='replace').read()
                           for u in compiled if u.startswith('platform/'))
    menudef = zipfile.ZipFile(os.path.join(ROMFS, 'zdoom.pk3')).read('menudef.txt').decode('latin1')
    engine_menudef = open(os.path.join(ROOT, 'engine', 'zdoom', 'wadsrc', 'static', 'menudef.txt'), encoding='latin1').read()
    out = subprocess.run([SIZE, ELF], capture_output=True, text=True).stdout.split('\n')[1].split()
    binary_mb = (int(out[0]) + int(out[1]) + int(out[2])) / 2**20
    has = lambda s: s in header
    uses = lambda f: ('platform/3ds/' + f) in compiled

    def check_sources(r):
        return r['unit'] in compiled if r.get('compile') == 'yes' else r.get('reason', '~') != '~'

    mods_ok = {}
    for s in sources:
        if s.get('compile') == 'yes':
            mods_ok.setdefault(s['module'], True)
            mods_ok[s['module']] &= s['unit'] in compiled

    # CIA rows: build/doomlauncher3ds.cia is newer than the ELF and its NCCH carries the sheet's title ID
    B = {r['key']: str(r.get('value', '~')) for r in rows('build', 'Build')}
    cia_ok = False
    if os.path.exists(CIA) and os.path.getmtime(CIA) >= os.path.getmtime(ELF):
        data = open(CIA, 'rb').read(0x10000)
        tid = struct.pack('<Q', 0x0004000000000000 | (int(B['unique_id'], 16) << 8))
        cia_ok = tid in data or tid[::-1] in data  # NCCH header little-endian, TMD big-endian
        # exheader kernel caps must map DSP memory (0x1FF00000) or ndspInitialize data-aborts on hardware
        cia_ok = cia_ok and struct.pack('<I', 0xFF81FF00) in data

    def check_build(r):
        k, v = r['key'], str(r.get('value', '~'))
        if k in ('arch_flags', 'libs', 'ld_extra'):
            return all(t in makefile for t in v.split())
        if k == 'icon_source':
            return os.path.exists(os.path.join(ROOT, 'platform', '3ds', 'icon.png'))
        if k == 'engine_binary_mb':
            return float(v) >= binary_mb
        if k in CIA_KEYS:
            return cia_ok
        return has('#define BD3DS_%s ' % re.sub(r'\W', '_', k).upper())

    def check_subsystem(r):
        impl = r.get('impl', '~')
        if impl == '~':
            return r.get('notes', '~') != '~'	# intentionally not built (crash catcher, x86 asm)
        if impl.endswith(('.cpp', '.c')):
            return impl in compiled
        d = impl.split(' (')[0]
        return any(u.startswith(d) for u in compiled)

    def check_patch(r):  # every edited engine file carries a [3DS] marker and its object is up to date
        files = [f.strip() for f in r['file'].split('+')]
        base = os.path.dirname(files[0])
        for f in (f if '/' in f else base + '/' + f for f in files):  # "dir/a.cpp + a.h"
            src = os.path.join(ROOT, f)
            if not os.path.exists(src) or '[3DS]' not in open(src, encoding='utf-8', errors='replace').read():
                return False
            if f.endswith(('.cpp', '.c')) and f not in compiled:  # headers: marker only
                return False
        return True

    checks = {
        ('build', 'Sources'): check_sources,
        ('build', 'Patches'): check_patch,
        ('build', 'Modules'): lambda r: mods_ok.get(r['module'], False),
        ('build', 'Build'): check_build,
        ('platform', 'Subsystems'): check_subsystem,
        ('platform', 'Video'): lambda r: has('"%s"' % r['menu_label']) and uses('video_3ds.cpp') and uses('hardware_3ds.cpp'),
        ('platform', 'Controls'): lambda r: has('"%s"' % r['button']) and uses('i_input_3ds.cpp') and uses('i_joystick_3ds.cpp'),
        ('platform', 'Menu'): lambda r: ('"%s", "%s"' % (r['option'], r['cvar']) in menudef and r['cvar'] in platform_src)
                                        if r.get('cvar', '~') != '~' else 'CustomizeControls' in engine_menudef,
        ('platform', 'BottomScreen'): lambda r: has('"%s"' % r['element']) and uses('bottom_3ds.cpp'),
        ('platform', 'BottomModes'): lambda r: has('{ "%s", "%s"' % (r['mode'], r['menu_label'])) and uses('bottom_3ds.cpp'),
        ('platform', 'Launcher'): lambda r: has('{ "%s", "%s", "%s"' % (r['element'], r['kind'], r['screen'])) and uses('launcher_3ds.cpp'),
        ('platform', 'Compat'): lambda r: r['resolution'] != 'implemented' or all(
            'Compat sheet' in (open(os.path.join(ROOT, 'engine', 'zdoom', 'src', f) if os.path.exists(os.path.join(ROOT, 'engine', 'zdoom', 'src', f))
                               else os.path.join(ROOT, 'engine', 'zdoom', 'src', 'thingdef', f), encoding='latin1').read())
            for f in r['files'].split()),
        ('platform', 'ModCache'): lambda r: has('#define BD3DS_MODCACHE_%s ' % r['rule'].upper()) and uses('modcache_3ds.cpp')
                                         and ('BD3DS_MODCACHE_%s' % r['rule'].upper() in platform_src or r['rule'].startswith('cache_')),
        ('platform', 'Audio'): lambda r: has('BD3DS_AUDIO_%s ' % re.sub(r'\W', '_', r['item']).upper()) and uses('ndspsound.cpp') and uses('vorbis_decoder.cpp'),
        ('platform', 'Filesystem'): lambda r: os.path.exists(os.path.join(ROMFS, r['path'][7:])) if r['path'].startswith('romfs:/')
                                        else (has('BD3DS_PATH_%s ' % r['role'].upper()) and (uses('i_specialpaths_3ds.cpp') or uses('i_system_3ds.cpp'))),
    }
    summary = {}
    for book in ('build', 'platform'):
        path = os.path.join(SHEETS, book + '.xlsx')
        wb = load_workbook(path)
        for (b, sheet), fn in checks.items():
            if b != book:
                continue
            ws = wb[sheet]
            hdr = [c.value for c in ws[1]]
            si = hdr.index('status') + 1
            data = rows(book, sheet)
            done = 0
            for i, r in enumerate(data):
                ok = bool(fn(r))
                ws.cell(row=i + 2, column=si, value='done' if ok else 'todo')
                done += ok
            summary[book + '/' + sheet] = '%d/%d' % (done, len(data))
        wb.save(path)
    print('binary %.1f MB' % binary_mb, summary)

if __name__ == '__main__':
    main()
